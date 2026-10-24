"""Reading uploaded spreadsheets (.xlsx or .csv) into header-keyed rows, and writing import
templates (PRD F16). Knows nothing about people or permissions — that is bulk_import.py."""

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from itertools import islice
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from app.errors import ApiError

MAX_ROWS = 2000
XLSX_MAGIC = b"PK\x03\x04"
DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y")  # Indian sheets: day first


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    required: bool
    help: str
    example: str


@dataclass(frozen=True)
class Row:
    number: int  # spreadsheet row number, as the admin sees it (header = 1)
    values: dict[str, str | date | None]


def _normalise_header(value: Any) -> str:
    return " ".join(str(value or "").replace("*", "").lower().split())


def _cell(value: Any) -> str | date | None:
    """Cells arrive as str / int / float / datetime; keep dates, stringify the rest. A whole
    float (roll number 2024001 stored as a number) loses its '.0'."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _xlsx_rows(content: bytes) -> list[tuple[Any, ...]]:
    # ponytail: openpyxl parses XML with the stdlib expat (billion-laughs protected since
    # expat 2.4); the upload is size-capped. Add defusedxml if files ever come from outside.
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return list(islice(book.worksheets[0].iter_rows(values_only=True), MAX_ROWS + 2))
        finally:
            book.close()
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ApiError(422, "unreadable_file", "This file could not be read as Excel.") from exc


def _csv_rows(content: bytes) -> list[tuple[Any, ...]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")  # Excel on Windows saves "CSV" in the ANSI code page
    return [tuple(r) for r in islice(csv.reader(io.StringIO(text)), MAX_ROWS + 2)]


def read_rows(content: bytes, columns: Sequence[Column]) -> list[Row]:
    """Rows keyed by column key. Unknown headers are ignored; missing required headers are a
    file-level error. Blank rows are skipped."""
    raw = _xlsx_rows(content) if content.startswith(XLSX_MAGIC) else _csv_rows(content)
    if not raw:
        raise ApiError(422, "empty_file", "The file is empty.")
    by_header = {_normalise_header(c.header): c.key for c in columns}
    positions = {
        by_header[_normalise_header(h)]: i
        for i, h in enumerate(raw[0])
        if _normalise_header(h) in by_header
    }
    missing = [c.header for c in columns if c.required and c.key not in positions]
    if missing:
        raise ApiError(
            422,
            "missing_columns",
            "The first row must name the columns. Missing: " + ", ".join(missing) + ".",
            {"missing": missing},
        )
    if len(raw) > MAX_ROWS + 1:
        raise ApiError(422, "too_many_rows", f"Import at most {MAX_ROWS} rows per file.")
    rows = []
    for number, values in enumerate(raw[1:], start=2):
        cells = {key: _cell(values[i]) if i < len(values) else None for key, i in positions.items()}
        if any(v is not None for v in cells.values()):
            rows.append(Row(number, cells))
    return rows


def parse_date(value: str | date | None) -> date | None:
    if value is None or isinstance(value, date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError("use a date like 2024-07-15 or 15/07/2024")


def template(columns: Sequence[Column], title: str) -> bytes:
    """Sheet 1: just the header row, ready to fill. Sheet 2: what each column means."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = title
    sheet.append([c.header + (" *" if c.required else "") for c in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for i, c in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = max(16, len(c.header) + 4)
    guide = book.create_sheet("How to fill")
    guide.append(["Column", "Required", "What to enter", "Example"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for c in columns:
        guide.append([c.header, "yes" if c.required else "no", c.help, c.example])
    guide.append([])
    guide.append(["Only the first sheet is imported. Delete no header; extra columns are ignored."])
    for letter, width in (("A", 22), ("B", 10), ("C", 70), ("D", 28)):
        guide.column_dimensions[letter].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
