"""Bulk import (PRD F16): dry run, all-or-nothing, same rules as the forms, safe re-runs."""

import dataclasses
import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app import notify
from app.config import Smtp, settings
from app.models import Faculty, Milestone, Notification, Scholar, User
from tests.conftest import auth

SMTP = Smtp(host="smtp.test", port=587, user="u", password="p", mail_from="noreply@x.edu")


def csv_bytes(*lines: str) -> bytes:
    return ("\n".join(lines) + "\n").encode()


def xlsx_bytes(rows: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def upload(client, user, kind, content, **params):
    params.setdefault("dry_run", "false")
    return await client.post(
        f"/api/v1/import/{kind}", headers=auth(user.id), content=content, params=params
    )


async def count(session, model) -> int:
    return await session.scalar(
        select(func.count()).select_from(model).where(model.deleted_at.is_(None))
    )


FACULTY_CSV = csv_bytes(
    "Name *,Email *,Designation *,Unit *,Max scholars",
    "Dr. One,one@x.edu,Professor,A1,",
    "Dr. Two,two@x.edu,assistant professor,Dept A1,3",
)


async def test_template_downloads_for_admins_only(client, world):
    res = await client.get("/api/v1/import/scholars/template", headers=auth(world.admin_a1.id))
    assert res.status_code == 200
    book = load_workbook(io.BytesIO(res.content))
    headers = [c.value for c in book.worksheets[0][1]]
    assert headers[:3] == ["Roll number *", "Name *", "Email *"]
    assert book.sheetnames == ["Scholars", "How to fill"]
    denied = await client.get("/api/v1/import/scholars/template", headers=auth(world.scholar_a1.id))
    assert denied.status_code == 403


async def test_dry_run_writes_nothing_then_import_creates(client, session, world):
    before = await count(session, Faculty)
    dry = await upload(client, world.admin_a1, "faculty", FACULTY_CSV, dry_run="true")
    assert dry.status_code == 200
    assert dry.json()["to_create"] == 2 and dry.json()["errors"] == []
    assert await count(session, Faculty) == before

    real = await upload(client, world.admin_a1, "faculty", FACULTY_CSV)
    assert real.json()["created"] == 2
    two = await session.scalar(
        select(Faculty).join(User, User.id == Faculty.user_id).where(User.email == "two@x.edu")
    )
    assert two.designation == "Assistant Professor" and two.max_scholars == 3
    one = await session.scalar(select(User).where(User.email == "one@x.edu"))
    assert one.password_hash is None  # logs in via link or SSO only


async def test_rerun_skips_rows_already_present(client, session, world):
    await upload(client, world.admin_a1, "faculty", FACULTY_CSV)
    again = await upload(client, world.admin_a1, "faculty", FACULTY_CSV)
    body = again.json()
    assert body["created"] == 0 and body["errors"] == []
    assert [s["row"] for s in body["already_present"]] == [2, 3]


async def test_any_bad_row_means_nothing_is_written(client, session, world):
    before = await count(session, User)
    content = csv_bytes(
        "Name,Email,Designation,Unit,Max scholars",
        "Dr. Good,good@x.edu,Professor,A1,",
        "Dr. Bad,not-an-email,Lecturer,NOPE,",
        "Dr. Greedy,greedy@x.edu,Assistant Professor,A1,9",
        "Dr. Dup,good@x.edu,Professor,A1,",
    )
    res = await upload(client, world.admin_a1, "faculty", content)
    errors = {(e["row"], e["column"]) for e in res.json()["errors"]}
    assert errors == {
        (3, "Email"),
        (3, "Designation"),
        (3, "Unit"),
        (4, "Max scholars"),
        (5, "Email"),
    }
    assert res.json()["created"] == 0
    assert await count(session, User) == before


async def test_units_outside_the_admins_subtree_are_row_errors(client, world):
    content = csv_bytes("Name,Email,Designation,Unit", "Dr. B,b@x.edu,Professor,B1")
    res = await upload(client, world.admin_a1, "faculty", content)
    assert res.json()["errors"][0]["message"] == "Dept B1 is outside the units you manage"
    # The root admin may import anywhere.
    ok = await upload(client, world.root_admin, "faculty", content)
    assert ok.json()["created"] == 1


async def test_only_admins_may_import(client, world):
    for user in (world.hod_a2, world.research_office, world.scholar_a1, world.sup_a2):
        res = await client.post(
            "/api/v1/import/faculty", headers=auth(user.id), content=FACULTY_CSV
        )
        assert res.status_code == 403, user


async def test_scholars_from_excel_with_supervisor_under_cap(client, session, world):
    world.meera_faculty.max_scholars = 2  # meera already guides scholar_a1
    session.add(world.meera_faculty)
    await session.commit()
    content = xlsx_bytes(
        [
            [
                "Roll number *",
                "Name *",
                "Email *",
                "Unit *",
                "Supervisor email",
                "Stage",
                "Admission date",
            ],
            [
                2024001,
                "Asha",
                "asha@x.edu",
                "A1",
                "meera@x.edu",
                "Coursework",
                datetime(2024, 7, 15),
            ],
            [2024002, "Ravi", "ravi@x.edu", "A1", "MEERA@x.edu", None, "16/07/2024"],
        ]
    )
    res = await upload(client, world.admin_a1, "scholars", content)
    body = res.json()
    assert body["errors"] == [
        {
            "row": 3,
            "column": "Supervisor email",
            "message": "already guides 2 scholars — the UGC limit is 2",
        }
    ]
    assert body["created"] == 0

    # Without Ravi's supervisor the file is clean.
    content = xlsx_bytes(
        [
            ["Roll number", "Name", "Email", "Unit", "Supervisor email", "Stage", "Admission date"],
            [
                2024001,
                "Asha",
                "asha@x.edu",
                "A1",
                "meera@x.edu",
                "Coursework",
                datetime(2024, 7, 15),
            ],
            [2024002, "Ravi", "ravi@x.edu", "A1", None, None, "16/07/2024"],
        ]
    )
    res = await upload(client, world.admin_a1, "scholars", content)
    assert res.json()["created"] == 2
    asha = await session.scalar(select(Scholar).where(Scholar.roll_number == "2024001"))
    assert asha.supervisor_id == world.meera_faculty.id
    assert asha.stage == "Coursework" and asha.admission_date == date(2024, 7, 15)
    ravi = await session.scalar(select(Scholar).where(Scholar.roll_number == "2024002"))
    assert ravi.admission_date == date(2024, 7, 16) and ravi.stage == "Registered"
    # Importing history is quiet: no "now your scholar" notification per row.
    assert await session.scalar(select(func.count()).select_from(Notification)) == 0


async def test_scholar_row_checks(client, world):
    content = csv_bytes(
        "Roll number,Name,Email,Unit,Supervisor email,Co-supervisor email,Admission date",
        "R1,One,r1@x.edu,A1,nobody@x.edu,,",
        "R2,Two,r2@x.edu,A1,,co_sup@x.edu,31/02/2024",
        "PHD-A1-001,Existing,fresh@x.edu,A1,,,",
    )
    res = await upload(client, world.admin_a1, "scholars", content, dry_run="true")
    body = res.json()
    messages = {(e["row"], e["column"]): e["message"] for e in body["errors"]}
    assert messages[(2, "Supervisor email")] == "no faculty member with email nobody@x.edu"
    assert "switched off" in messages[(3, "Co-supervisor email")]
    assert "date" in messages[(3, "Admission date")]
    assert body["already_present"] == [{"row": 4, "what": "PHD-A1-001"}]


async def test_history_import_and_rerun(client, session, world):
    content = csv_bytes(
        "Roll number,Milestone,Status,Approval body,Resolution number,Resolution date",
        "PHD-A1-001,Proposal approval,,rdc,RDC/2023/14,22/09/2023",
        "PHD-A1-001,proposal_approval,,RDC,RDC/2023/14,2023-09-22",
        "PHD-B1-001,coursework,,RDC,RDC/2023/15,2023-09-22",
        "NOPE,coursework,,RDC,,",
    )
    res = await upload(client, world.admin_a1, "history", content, dry_run="true")
    errors = {(e["row"], e["column"]) for e in res.json()["errors"]}
    assert errors == {(4, "Roll number"), (5, "Roll number")}  # B1 out of scope, unknown

    clean = csv_bytes(*content.decode().splitlines()[:3])
    res = await upload(client, world.admin_a1, "history", clean)
    assert res.json()["created"] == 1
    assert res.json()["already_present"] == [{"row": 3, "what": "PHD-A1-001 proposal_approval"}]
    m = await session.scalar(select(Milestone).where(Milestone.scholar_id == world.s_a1.id))
    assert (m.status, m.approval_body, m.resolution_date) == ("Approved", "RDC", date(2023, 9, 22))
    again = await upload(client, world.admin_a1, "history", clean)
    assert again.json()["created"] == 0


async def test_invitations_go_out_after_import(client, world, monkeypatch):
    monkeypatch.setattr(notify, "settings", dataclasses.replace(settings, smtp=SMTP))
    sent = []

    async def fake_send(msg):
        sent.append(msg["To"])

    monkeypatch.setattr(notify, "_send", fake_send)
    res = await upload(client, world.admin_a1, "faculty", FACULTY_CSV, send_invitations="true")
    assert res.json()["invitations_sent"] == 2
    assert sorted(sent) == ["one@x.edu", "two@x.edu"]


async def test_invitations_need_email(client, world):
    res = await upload(client, world.admin_a1, "faculty", FACULTY_CSV, send_invitations="true")
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "email_not_configured"


@pytest.mark.parametrize(
    ("content", "code"),
    [
        (b"", "empty_file"),
        (csv_bytes("Name,Email", "x,y@x.edu"), "missing_columns"),
        (b"PK\x03\x04not really a zip", "unreadable_file"),
    ],
)
async def test_unusable_files(client, world, content, code):
    res = await upload(client, world.admin_a1, "faculty", content)
    assert res.status_code == 422
    assert res.json()["error"]["code"] == code


async def test_oversized_file_is_refused(client, world):
    res = await upload(client, world.admin_a1, "faculty", b"x" * (5 * 1024 * 1024 + 1))
    assert res.status_code == 413
